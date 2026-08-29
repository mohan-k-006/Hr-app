import os
import pymupdf as fitz
import docx
import logging

logger = logging.getLogger(__name__)

def extract_text_from_file(file_path: str, filename: str) -> tuple[str, str]:
    """
    Extracts raw text content from PDF, DOCX, or TXT files.
    Returns tuple of (raw_text, file_type)
    """
    ext = os.path.splitext(filename)[1].lower().strip(".")
    
    if ext == "pdf":
        return extract_text_from_pdf(file_path), "pdf"
    elif ext == "docx":
        return extract_text_from_docx(file_path), "docx"
    elif ext in ["txt", "md"]:
        return extract_text_from_txt(file_path), "txt"
    else:
        raise ValueError(f"Unsupported file format: .{ext}. Supported formats are PDF, DOCX, TXT.")

def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    try:
        doc = fitz.open(file_path)
        for page in doc:
            page_text = page.get_text()
            if page_text:
                text += page_text + "\n"
        doc.close()
    except Exception as e:
        logger.error(f"Error reading PDF file {file_path}: {e}")
        raise ValueError(f"Could not extract text from PDF: {str(e)}")
    return text.strip()

def extract_text_from_docx(file_path: str) -> str:
    text = ""
    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            if p.text:
                text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text += row_text + "\n"
    except Exception as e:
        logger.error(f"Error reading DOCX file {file_path}: {e}")
        raise ValueError(f"Could not extract text from DOCX: {str(e)}")
    return text.strip()

def extract_text_from_txt(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception as e:
        logger.error(f"Error reading TXT file {file_path}: {e}")
        raise ValueError(f"Could not extract text from TXT file: {str(e)}")
